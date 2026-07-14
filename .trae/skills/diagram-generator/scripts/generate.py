#!/usr/bin/env python3
"""
PlantUML 图表生成脚本
功能:
  1. 将 PlantUML 代码编码并从 PlantUML 服务器渲染为 PNG
  2. (可选) 将 PNG 打包进 Word 文档
  3. (可选) 在浏览器打开在线编辑器

用法:
  # 渲染为 PNG(PKB 驱动模式常用)
  python generate.py --code @diagram.puml --output diagram.png --format png --no-browser

  # 渲染并打包为 DOCX(自然语言模式)
  python generate.py --request "登录流程" --code @diagram.puml --output report.docx

  # 内联代码
  python generate.py --code "@startuml\nBob -> Alice : hi\n@enduml" --output test.png --format png
"""

import argparse
import logging
import os
import sys
import urllib.request
import urllib.error
import webbrowser
from pathlib import Path

# PlantUML 服务器(主 + 备用)
PLANTUML_SERVERS = [
    "https://www.plantuml.com/plantuml",
    "https://plantuml.com/plantuml",
]

logging.basicConfig(
    level=logging.INFO,
    format="%(levelname)s %(message)s",
)
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except AttributeError:
    pass

# PlantUML 使用的特殊编码字符表
PLANTUML_ENCODE_DICT = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz-_"


def _encode6bit(b: int) -> str:
    """编码单个6位值为PlantUML字符"""
    return PLANTUML_ENCODE_DICT[b]


def _append3bytes(b1: int, b2: int, b3: int) -> str:
    """将3字节编码为4个PlantUML字符"""
    c1 = b1 >> 2
    c2 = ((b1 & 0x3) << 4) | (b2 >> 4)
    c3 = ((b2 & 0xF) << 2) | (b3 >> 6)
    c4 = b3 & 0x3F
    return _encode6bit(c1) + _encode6bit(c2) + _encode6bit(c3) + _encode6bit(c4)


def encode_plantuml(text: str) -> str:
    """
    将 PlantUML 文本编码为 PlantUML 服务器 URL 使用的格式。
    算法:UTF-8 字节 → raw DEFLATE 压缩(无zlib头) → PlantUML base64 编码
    """
    import zlib
    raw = text.encode("utf-8")
    # PlantUML 服务器期望 raw DEFLATE(RFC 1951),不是 zlib(RFC 1950)
    # wbits=-15 表示 raw deflate,不包含 zlib 头和尾部校验
    compress_obj = zlib.compressobj(9, zlib.DEFLATED, -15)
    compressed = compress_obj.compress(raw) + compress_obj.flush()

    # 每次取3字节进行编码
    result = []
    for i in range(0, len(compressed), 3):
        b1 = compressed[i]
        b2 = compressed[i + 1] if i + 1 < len(compressed) else 0
        b3 = compressed[i + 2] if i + 2 < len(compressed) else 0
        result.append(_append3bytes(b1, b2, b3))
    return "".join(result)


def render_png(plantuml_code: str, servers=None) -> bytes:
    """
    将 PlantUML 代码渲染为 PNG 图片字节。
    依次尝试多个服务器,全部失败则抛出异常。
    """
    if servers is None:
        servers = PLANTUML_SERVERS

    encoded = encode_plantuml(plantuml_code)
    last_error = None

    for server in servers:
        url = f"{server}/png/{encoded}"
        try:
            req = urllib.request.Request(url, headers={
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36"
            })
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = resp.read()
                # 校验:必须是真正的 PNG 文件
                if data[:4] == b"\x89PNG":
                    return data
                last_error = f"服务器 {server} 返回非PNG数据(可能是编码错误)"
        except Exception as e:
            last_error = f"服务器 {server} 请求失败: {e}"
            continue

    raise RuntimeError(f"所有 PlantUML 服务器均渲染失败。最后错误: {last_error}")


def save_png(png_data: bytes, output_path: str):
    """保存 PNG 数据到文件"""
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output_path + ".tmp"
    with open(tmp_path, "wb") as f:
        f.write(png_data)
    os.replace(tmp_path, output_path)


def create_docx(plantuml_code: str, png_data: bytes, request_text: str, output_path: str):
    """
    将图表打包为 Word 文档,包含:
    - 用户请求描述
    - PlantUML 源代码
    - 图表 URL
    - 渲染后的图片
    """
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError:
        # python-docx 未安装,退化为只保存 PNG
        png_path = output_path.rsplit(".", 1)[0] + ".png"
        save_png(png_data, png_path)
        logging.warning(f"python-docx 未安装,已保存 PNG: {png_path}")
        return

    import io
    doc = Document()

    # 标题
    doc.add_heading("PlantUML 图表报告", level=0)

    # 用户请求
    doc.add_heading("需求描述", level=1)
    doc.add_paragraph(request_text)

    # PlantUML 代码
    doc.add_heading("PlantUML 源代码", level=1)
    code_para = doc.add_paragraph()
    run = code_para.add_run(plantuml_code)
    run.font.name = "Consolas"
    run.font.size = doc.styles["Normal"].font.size

    # 图表 URL
    doc.add_heading("在线编辑链接", level=1)
    encoded = encode_plantuml(plantuml_code)
    url = f"https://www.plantuml.com/plantuml/uml/{encoded}"
    doc.add_paragraph(url)

    # 渲染图片
    doc.add_heading("图表预览", level=1)
    image_stream = io.BytesIO(png_data)
    doc.add_picture(image_stream, width=Inches(6))

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output_path + ".tmp"
    doc.save(tmp_path)
    os.replace(tmp_path, output_path)


def open_editor(plantuml_code: str):
    """在浏览器打开 PlantUML 在线编辑器"""
    encoded = encode_plantuml(plantuml_code)
    url = f"https://www.plantuml.com/plantuml/uml/{encoded}"
    webbrowser.open(url)
    return url


def main():
    parser = argparse.ArgumentParser(
        description="PlantUML 图表生成器",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  # 渲染 PNG
  python generate.py --code @diagram.puml --output er.png --format png --no-browser

  # 渲染并打包 DOCX
  python generate.py --request "ER图" --code @diagram.puml --output report.docx

  # 内联代码
  python generate.py --code "@startuml\\nA -> B\\n@enduml" --output test.png --format png
        """,
    )
    parser.add_argument("--request", default="", help="用户的自然语言需求描述")
    parser.add_argument(
        "--code",
        required=True,
        help='PlantUML 代码。用 @filename 从文件读取,或直接传入代码字符串',
    )
    parser.add_argument("--output", default="plantuml_output.png", help="输出文件路径")
    parser.add_argument(
        "--format",
        choices=["png", "docx"],
        default="docx",
        help="输出格式(默认 docx)",
    )
    parser.add_argument("--no-browser", action="store_true", help="不打开浏览器")
    parser.add_argument(
        "--server",
        default=None,
        help="自定义 PlantUML 服务器 URL",
    )

    args = parser.parse_args()

    # 读取 PlantUML 代码
    if args.code.startswith("@"):
        code_file = args.code[1:]
        if not os.path.exists(code_file):
            logging.error(f"错误: 文件不存在: {code_file}")
            sys.exit(1)
        with open(code_file, "r", encoding="utf-8") as f:
            plantuml_code = f.read()
    else:
        plantuml_code = args.code

    # 确保代码包含 @startuml / @enduml
    if "@start" not in plantuml_code:
        plantuml_code = "@startuml\n" + plantuml_code + "\n@enduml"

    # 配置服务器
    servers = [args.server] if args.server else PLANTUML_SERVERS

    # 公网服务器数据外发警告
    public_servers = {"https://www.plantuml.com/plantuml", "https://plantuml.com/plantuml"}
    if any(s in public_servers for s in servers):
        logging.warning(
            "⚠️ 安全提示:图表内容将通过公网发送到 PlantUML 服务器(含表名/API路径等)。\n"
            "   敏感项目请使用 --server http://localhost:8080 指定本地实例。",
        )

    logging.info("正在渲染 PlantUML 图表...")
    try:
        png_data = render_png(plantuml_code, servers)
    except RuntimeError as e:
        logging.error(f"渲染失败: {e}")
        sys.exit(1)

    # 根据格式输出
    if args.format == "png" or args.output.endswith(".png"):
        save_png(png_data, args.output)
        logging.info(f"PNG 已保存: {args.output}")
    else:
        create_docx(plantuml_code, png_data, args.request, args.output)
        logging.info(f"DOCX 已保存: {args.output}")

    # 打开浏览器
    if not args.no_browser:
        url = open_editor(plantuml_code)
        logging.info(f"在线编辑器已打开: {url}")

    logging.info("完成!")


if __name__ == "__main__":
    main()
